"""Build 01-Business-Strategy/Customer-Persona-Builder.xlsx (AQVANI.SHOP · AI Business OS 2026).

Run:  python3 build_persona_builder.py
Then recalculate with LibreOffice (see xlsx skill recalc.py) so cached values exist.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[2] / "01-Business-Strategy" / "Customer-Persona-Builder.xlsx"

# Brand
DARK, SECOND, GOLD, CREAM, CHAR, WHITE = "073B2A", "0B4D35", "D9A51A", "F8F1E7", "17231E", "FFFFFF"
CALC = "E6EFEA"  # light green tint for formula cells
FONT = "Arial"

thin = Side(style="thin", color="C9C2B6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def f(size=10, bold=False, color=CHAR, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", start_color=c, end_color=c)


def title(ws, text, sub, width_cols):
    ws["A1"] = text
    ws["A1"].font = f(16, True, WHITE)
    ws["A2"] = sub
    ws["A2"].font = f(10, False, CREAM, italic=True)
    for r in (1, 2):
        for c in range(1, width_cols + 1):
            ws.cell(row=r, column=c).fill = fill(DARK)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18


def header(ws, row, labels, start_col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=lab)
        c.font = f(10, True, WHITE)
        c.fill = fill(SECOND)
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 32


def inp(cell, value=None):
    if value is not None:
        cell.value = value
    cell.fill = fill(CREAM)
    cell.font = f()
    cell.border = BORDER
    cell.alignment = WRAP


def calc(cell, formula, fmt=None):
    cell.value = formula
    cell.fill = fill(CALC)
    cell.font = f(10, True)
    cell.border = BORDER
    cell.alignment = CENTER
    if fmt:
        cell.number_format = fmt


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def page(ws, landscape=True):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "AQVANI.SHOP · AI Business Operating System 2026 · Customer Persona Builder"


wb = Workbook()

# ---------------------------------------------------------------- 1. Start Here
ws = wb.active
ws.title = "Start Here"
title(ws, "CUSTOMER PERSONA BUILDER", "AI Business Operating System 2026 · System 01 — Business Strategy · AQVANI.SHOP", 4)
widths(ws, {"A": 4, "B": 30, "C": 70, "D": 4})
rows = [
    ("What this workbook does", "Builds evidence-based customer personas, ranks customer pain points, scores prospects against your Ideal Customer Profile (ICP), and keeps a log of real customer conversations. It supports Workbook Tools 2.2–2.4 (worksheets W-STR-04, 05, 06)."),
    ("Step 1 — Interview Log", "Record each real customer conversation (use prompt P-STR-009 for questions). Use codes, not names or phone numbers."),
    ("Step 2 — Persona Builder", "Fill up to three personas using evidence from the Interview Log, reviews and enquiries. Completion % updates automatically. Use prompt P-STR-008 to sharpen."),
    ("Step 3 — Pain Point Scoring", "List pains in customer language and rate four factors 1–5. Score and rank calculate automatically. Use prompts P-STR-011 and P-STR-012."),
    ("Step 4 — ICP Scorecard", "Enter 5–7 fit criteria and weights (must total 100%). Then score each new prospect 0–5 per criterion to get a Fit % and band. Use prompt P-STR-010."),
    ("Step 5 — Carry forward", "Copy your top 3 pains, persona one-liner and ICP sentence into W-STR-22 (One-Page Strategy Summary)."),
    ("Colour legend", "CREAM cells = your input (type here) · GREEN-TINT cells = automatic formulas (do not type here) · DARK GREEN = headers."),
    ("Example data", "Rows marked 'Example' use fictional businesses (BrightNest Cleaning Co. and Northbridge Ops Advisory). They show the expected format. Overwrite or delete them when you start."),
    ("Formulas used", "Pain Score = SUMPRODUCT(ratings, weights) ÷ SUM(weights) × 20 (range 20–100). ICP Fit % = SUMPRODUCT(ratings, weights) ÷ 5. Completion % = COUNTA(answers) ÷ number of fields."),
    ("Google Sheets", "Upload this file to Google Drive and open with Google Sheets. All formulas and dropdowns are compatible."),
    ("Privacy", "Do not store customer phone numbers, addresses or ID details here. Use codes such as C-014."),
    ("AI review rule", "AI can organise your evidence; it cannot replace it. Review every AI output before using it in marketing or sales."),
]
r = 4
for k, v in rows:
    a = ws.cell(row=r, column=2, value=k)
    a.font = f(10, True, DARK)
    a.alignment = WRAP
    b = ws.cell(row=r, column=3, value=v)
    b.font = f()
    b.alignment = WRAP
    ws.row_dimensions[r].height = 46
    r += 1
ws.cell(row=r + 1, column=2, value="Fictional examples only. Not customer results. General guidance, not legal or financial advice.").font = f(9, italic=True)
page(ws, landscape=False)

# ---------------------------------------------------------------- 6. Lists (built early; referenced by validation)
lists = wb.create_sheet("Lists")

# ---------------------------------------------------------------- 2. Persona Builder
pb = wb.create_sheet("Persona Builder", 1)
title(pb, "PERSONA BUILDER  ·  W-STR-04", "Up to three personas side by side. Cream = input. Base each answer on evidence; mark guesses as INFERRED.", 6)
widths(pb, {"A": 30, "B": 42, "C": 40, "D": 40, "E": 40, "F": 2})
header(pb, 4, ["Section", "Guidance", "Persona 1 (Example — overwrite)", "Persona 2", "Persona 3"])
fields = [
    ("Persona name & one-line description", "A memorable name + who they are in one line"),
    ("Snapshot", "Age range, role, life/business stage, location"),
    ("Situation right now", "What is happening in their life or business"),
    ("Top 3 goals", "What they are trying to achieve"),
    ("Top 3 frustrations", "What gets in the way"),
    ("Trigger moments", "What makes them look for help now"),
    ("Current solution / alternative", "What they do today instead (incl. DIY / nothing)"),
    ("Top 3 objections", "Why they might hesitate to buy"),
    ("Decision process", "Who decides, how long, compared with what"),
    ("Channels", "Where they spend attention and ask for advice"),
    ("Their exact words", "3–5 real quotes from conversations, reviews or chats"),
    ("Evidence used", "Number and type of sources (e.g. 4 calls, 9 chats)"),
]
example = [
    "Busy Apartment Couple — 'Priya & Arjun': both work full-time and want a reliably clean home without supervising anyone",
    "30–42, salaried professionals, 2–3BHK apartment, Pune west, often with a young child",
    "Long workdays and commutes; part-time help is inconsistent; weekends spent on chores",
    "1. A consistently clean home  2. Weekends free  3. Help they can trust in the house",
    "1. Corners, fans and grout always missed  2. No-shows  3. Chasing and re-explaining every time",
    "Guests or family visiting; festivals; a no-show on a busy week; moving in",
    "Part-time maid + occasional deep clean before festivals",
    "1. 'Will the same people come each time?'  2. 'Is it safe to leave them alone at home?'  3. 'Monthly plan sounds expensive'",
    "Joint decision; compare with maid cost and one-off deep clean; decide within a week of a trigger",
    "Apartment WhatsApp groups, Instagram, Google search, neighbours' recommendations",
    "'They clean what's visible and leave the rest.' · 'I just don't want to follow up every week.' · 'Same team would be great.'",
    "4 customer calls, 9 enquiry chats, 6 competitor reviews (example)",
]
FIRST, LAST = 5, 5 + len(fields) - 1
for i, (sec, guide) in enumerate(fields):
    row = FIRST + i
    a = pb.cell(row=row, column=1, value=sec)
    a.font = f(10, True, DARK)
    a.border = BORDER
    a.alignment = WRAP
    g = pb.cell(row=row, column=2, value=guide)
    g.font = f(9, italic=True)
    g.border = BORDER
    g.alignment = WRAP
    inp(pb.cell(row=row, column=3), example[i])
    inp(pb.cell(row=row, column=4))
    inp(pb.cell(row=row, column=5))
    pb.row_dimensions[row].height = 48
status_row = LAST + 1
pb.cell(row=status_row, column=1, value="Persona status").font = f(10, True, DARK)
pb.cell(row=status_row, column=2, value="Draft → Evidence-based (5+ sources) → Validated (used in sales)").font = f(9, italic=True)
for col in (3, 4, 5):
    inp(pb.cell(row=status_row, column=col), "Evidence-based" if col == 3 else None)
comp_row = status_row + 1
pb.cell(row=comp_row, column=1, value="Completion %").font = f(10, True, DARK)
pb.cell(row=comp_row, column=2, value="Share of the 12 sections filled (automatic)").font = f(9, italic=True)
for col, L in ((3, "C"), (4, "D"), (5, "E")):
    calc(pb.cell(row=comp_row, column=col), f"=COUNTA({L}{FIRST}:{L}{LAST})/ROWS({L}{FIRST}:{L}{LAST})", "0%")
for r_ in (status_row, comp_row):
    for c_ in (1, 2):
        pb.cell(row=r_, column=c_).border = BORDER
dv_status = DataValidation(type="list", formula1="=Lists!$C$2:$C$4", allow_blank=True)
pb.add_data_validation(dv_status)
dv_status.add(f"C{status_row}:E{status_row}")
pb.conditional_formatting.add(
    f"C{comp_row}:E{comp_row}",
    ColorScaleRule(start_type="num", start_value=0, start_color="F4D6CC", end_type="num", end_value=1, end_color="B7D7C2"),
)
pb.freeze_panes = "C5"
page(pb)

# ---------------------------------------------------------------- 3. Pain Point Scoring
pp = wb.create_sheet("Pain Point Scoring", 2)
title(pp, "PAIN POINT SCORING  ·  W-STR-06",
      "Rate each factor 1–5. Score = weighted average × 20 (range 20–100). 75+ High · 55–74 Medium · below 55 Low.", 12)
widths(pp, {"A": 6, "B": 46, "C": 18, "D": 30, "E": 16, "F": 12, "G": 12, "H": 14, "I": 12, "J": 10, "K": 8, "L": 12})
pp["B3"] = "Factor weights (edit if needed — should total 100%)"
pp["B3"].font = f(10, True, DARK)
for col, lab in zip("CDEF", ["Frequency", "Intensity", "Willingness to pay", "Ability to solve"]):
    c = pp[f"{col}3"]
    c.value = lab
    c.font = f(9, True, WHITE)
    c.fill = fill(SECOND)
    c.alignment = CENTER
    c.border = BORDER
pp["G3"] = "Total"
pp["H3"] = "Check"
for col in "GH":
    pp[f"{col}3"].font = f(9, True, WHITE)
    pp[f"{col}3"].fill = fill(SECOND)
    pp[f"{col}3"].alignment = CENTER
    pp[f"{col}3"].border = BORDER
for col, w in zip("CDEF", [0.2, 0.3, 0.3, 0.2]):
    inp(pp[f"{col}4"], w)
    pp[f"{col}4"].number_format = "0%"
    pp[f"{col}4"].alignment = CENTER
calc(pp["G4"], "=SUM(C4:F4)", "0%")
calc(pp["H4"], '=IF(ROUND(G4,4)=1,"OK","Weights must total 100%")')
pp["C4"].comment = Comment("Default weights: Frequency 20%, Intensity 30%, Willingness to pay 30%, Ability to solve 20% (Aqvani default). Change to suit your market; the score divides by the weight total.", "Aqvani")
pp.row_dimensions[3].height = 30

H = 6
header(pp, H, ["#", "Pain point (customer language)", "Persona", "Source / evidence", "Source type",
               "Frequency 1–5", "Intensity 1–5", "Willingness to pay 1–5", "Ability to solve 1–5",
               "Score", "Rank", "Priority"])
P_FIRST, P_LAST = H + 1, H + 30
pain_examples = [
    ("The maid cleans the visible areas but corners, fans and grout are always dirty", "Persona 1", "9 enquiry chats", "Enquiry / chat", 4, 4, 4, 5),
    ("They confirm and then simply don't turn up", "Persona 1", "Competitor reviews", "Competitor review", 3, 5, 4, 4),
    ("I'm uncomfortable leaving strangers in my home when I'm at work", "Persona 1", "Calls with 4 customers", "Interview", 3, 5, 3, 4),
    ("Nobody is available in the two weeks before Diwali", "Persona 1", "Festive-season chats", "Enquiry / chat", 2, 4, 4, 3),
    ("Strong chemical smell — worried about kids and pets", "Persona 1", "3 reviews", "Own review", 2, 3, 3, 5),
    ("It takes forever to get a reply and a price", "Persona 1", "Own WhatsApp history", "Enquiry / chat", 3, 3, 2, 5),
]
for i in range(30):
    row = P_FIRST + i
    n = pp.cell(row=row, column=1, value=i + 1)
    n.font = f()
    n.alignment = CENTER
    n.border = BORDER
    for col in range(2, 10):
        inp(pp.cell(row=row, column=col))
    for col in range(6, 10):
        pp.cell(row=row, column=col).alignment = CENTER
    if i < len(pain_examples):
        ex = pain_examples[i]
        pp.cell(row=row, column=2, value=f"Example — {ex[0]}")
        for j, v in enumerate(ex[1:], start=3):
            pp.cell(row=row, column=j, value=v)
    calc(pp.cell(row=row, column=10),
         f'=IF(COUNT(F{row}:I{row})<4,"",ROUND(SUMPRODUCT(F{row}:I{row},$C$4:$F$4)/SUM($C$4:$F$4)*20,0))', "0")
    calc(pp.cell(row=row, column=11), f'=IF(J{row}="","",RANK(J{row},$J${P_FIRST}:$J${P_LAST}))', "0")
    calc(pp.cell(row=row, column=12), f'=IF(J{row}="","",IF(J{row}>=75,"High",IF(J{row}>=55,"Medium","Low")))')
    pp.row_dimensions[row].height = 30

dv_15 = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                       showErrorMessage=True, errorTitle="Rating 1–5", error="Enter a whole number from 1 to 5.")
pp.add_data_validation(dv_15)
dv_15.add(f"F{P_FIRST}:I{P_LAST}")
dv_w = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=False,
                      showErrorMessage=True, errorTitle="Weight", error="Enter a percentage between 0% and 100%.")
pp.add_data_validation(dv_w)
dv_w.add("C4:F4")
dv_persona = DataValidation(type="list", formula1="=Lists!$A$2:$A$4", allow_blank=True)
pp.add_data_validation(dv_persona)
dv_persona.add(f"C{P_FIRST}:C{P_LAST}")
dv_src = DataValidation(type="list", formula1="=Lists!$B$2:$B$9", allow_blank=True)
pp.add_data_validation(dv_src)
dv_src.add(f"E{P_FIRST}:E{P_LAST}")

pp.conditional_formatting.add(
    f"J{P_FIRST}:J{P_LAST}",
    ColorScaleRule(start_type="num", start_value=20, start_color="F4D6CC",
                   mid_type="num", mid_value=60, mid_color="FBE7B0",
                   end_type="num", end_value=100, end_color="B7D7C2"),
)
pp.conditional_formatting.add(f"L{P_FIRST}:L{P_LAST}",
                              CellIsRule(operator="equal", formula=['"High"'], fill=fill(DARK), font=Font(name=FONT, bold=True, color=WHITE)))
pp.conditional_formatting.add(f"L{P_FIRST}:L{P_LAST}",
                              CellIsRule(operator="equal", formula=['"Medium"'], fill=fill("FBE7B0")))
pp.conditional_formatting.add("H4", CellIsRule(operator="notEqual", formula=['"OK"'], fill=fill("F4D6CC")))
pp.cell(row=P_LAST + 2, column=2,
        value="Top 3 pains = rank 1–3. Copy them, in customer language, into W-STR-06 and W-STR-22.").font = f(9, italic=True)
pp.freeze_panes = f"C{P_FIRST}"
pp.auto_filter.ref = f"A{H}:L{P_LAST}"
page(pp)

# ---------------------------------------------------------------- 4. ICP Scorecard
icp = wb.create_sheet("ICP Scorecard", 3)
title(icp, "ICP SCORECARD  ·  W-STR-05",
      "Define 5–7 criteria with weights (total 100%). Score each prospect 0–5 per criterion. A red flag overrides the score.", 14)
widths(icp, {"A": 14, "B": 30, "C": 14, "D": 13, "E": 13, "F": 13, "G": 13, "H": 13, "I": 13, "J": 13,
             "K": 12, "L": 11, "M": 20, "N": 34})
icp["A3"] = "PART A — Define your Ideal Customer Profile"
icp["A3"].font = f(11, True, DARK)
header(icp, 4, ["#", "Fit criterion", "Weight %", "What '5' looks like", "", "", "What '0' looks like", "", ""])
icp.merge_cells("D4:F4")
icp.merge_cells("G4:I4")
C_FIRST, C_LAST = 5, 11
icp_examples = [
    ("Business type & size", 0.25, "Manufacturer, 20–150 staff, own production floor", "Trader or under 10 staff"),
    ("Pain urgency", 0.25, "Delays/rework hurting customer relationships now", "'Just exploring'"),
    ("Decision-maker access", 0.20, "Owner/MD in the first meeting", "Only a junior contact"),
    ("Budget readiness", 0.15, "Has paid consultants/software before", "Expects free advice"),
    ("Location", 0.15, "Within 150 km for site visits", "Frequent long-distance travel needed"),
]
for i in range(7):
    row = C_FIRST + i
    n = icp.cell(row=row, column=1, value=i + 1)
    n.font = f()
    n.alignment = CENTER
    n.border = BORDER
    inp(icp.cell(row=row, column=2))
    inp(icp.cell(row=row, column=3))
    icp.cell(row=row, column=3).number_format = "0%"
    icp.cell(row=row, column=3).alignment = CENTER
    icp.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
    icp.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
    inp(icp.cell(row=row, column=4))
    inp(icp.cell(row=row, column=7))
    for c_ in (5, 6, 8, 9):
        icp.cell(row=row, column=c_).border = BORDER
        icp.cell(row=row, column=c_).fill = fill(CREAM)
    if i < len(icp_examples):
        nm, w, five, zero = icp_examples[i]
        icp.cell(row=row, column=2, value=f"Example — {nm}")
        icp.cell(row=row, column=3, value=w)
        icp.cell(row=row, column=4, value=five)
        icp.cell(row=row, column=7, value=zero)
    icp.row_dimensions[row].height = 30
tot = C_LAST + 1
icp.cell(row=tot, column=2, value="Total weight (must be 100%)").font = f(10, True, DARK)
calc(icp.cell(row=tot, column=3), f"=SUM(C{C_FIRST}:C{C_LAST})", "0%")
calc(icp.cell(row=tot, column=4), f'=IF(ROUND(C{tot},4)=1,"OK","Weights must total 100%")')
icp.conditional_formatting.add(f"D{tot}", CellIsRule(operator="notEqual", formula=['"OK"'], fill=fill("F4D6CC")))
dv_wi = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
                       showErrorMessage=True, errorTitle="Weight", error="Enter a percentage between 0% and 100%.")
icp.add_data_validation(dv_wi)
dv_wi.add(f"C{C_FIRST}:C{C_LAST}")

icp.cell(row=tot + 2, column=1, value="Red flags (override the score):").font = f(10, True, DARK)
for k in range(3):
    rr = tot + 3 + k
    icp.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=9)
    inp(icp.cell(row=rr, column=2), ["Example — Wants guaranteed cost savings in writing",
                                     "Example — Will not share basic production data",
                                     "Example — Expects results without staff involvement"][k])

PB = tot + 8  # Part B title row
icp.cell(row=PB, column=1, value="PART B — Score prospects").font = f(11, True, DARK)
icp.cell(row=PB, column=4,
         value="Fit % = SUMPRODUCT(ratings, weights) ÷ 5 · 80–100% Ideal · 60–79% Good · 40–59% Possible · below 40% Poor").font = f(9, italic=True)
WROW = PB + 1  # helper weights row
HROW = PB + 2
icp.cell(row=WROW, column=3, value="Weight →").font = f(9, True, DARK)
crit_cols = "DEFGHIJ"
for i, col in enumerate(crit_cols):
    calc(icp[f"{col}{WROW}"], f"=IF($C${C_FIRST + i}=\"\",0,$C${C_FIRST + i})", "0%")
header(icp, HROW, ["Prospect code", "Notes (business / situation)", "Date"] + [""] * 7 +
       ["Red flag? (Yes/No)", "Fit %", "Band", "Recommended action"])
for i, col in enumerate(crit_cols):
    c = icp[f"{col}{HROW}"]
    c.value = f'=IF($B${C_FIRST + i}="","Criterion {i + 1}",$B${C_FIRST + i})'
icp.row_dimensions[HROW].height = 54
R_FIRST, R_LAST = HROW + 1, HROW + 40
prospect_examples = [
    ("Example P-001", "Auto-components unit, 60 staff; owner attended; delays hurting key account", (5, 5, 5, 4, 4), "No"),
    ("Example P-002", "Garment unit, 25 staff; manager only; 'exploring options'", (2, 3, 1, 2, 5), "No"),
    ("Example P-003", "Fabrication unit, 80 staff; wants savings guaranteed in contract", (5, 4, 5, 4, 4), "Yes"),
]
for i in range(40):
    row = R_FIRST + i
    inp(icp.cell(row=row, column=1))
    inp(icp.cell(row=row, column=2))
    inp(icp.cell(row=row, column=3))
    icp.cell(row=row, column=3).number_format = "dd-mmm-yyyy"
    for col in crit_cols:
        inp(icp[f"{col}{row}"])
        icp[f"{col}{row}"].alignment = CENTER
    inp(icp.cell(row=row, column=11))
    icp.cell(row=row, column=11).alignment = CENTER
    if i < len(prospect_examples):
        code, note, ratings, flag = prospect_examples[i]
        icp.cell(row=row, column=1, value=code)
        icp.cell(row=row, column=2, value=note)
        for col, v in zip(crit_cols, ratings):
            icp[f"{col}{row}"] = v
        icp.cell(row=row, column=11, value=flag)
    calc(icp.cell(row=row, column=12),
         f'=IF(COUNT(D{row}:J{row})=0,"",SUMPRODUCT(D{row}:J{row},$D${WROW}:$J${WROW})/5)', "0%")
    calc(icp.cell(row=row, column=13),
         f'=IF(L{row}="","",IF(K{row}="Yes","Red flag — review",IF(L{row}>=0.8,"Ideal",IF(L{row}>=0.6,"Good",IF(L{row}>=0.4,"Possible","Poor")))))')
    calc(icp.cell(row=row, column=14),
         f'=IF(M{row}="","",IF(M{row}="Ideal","Prioritise — respond same day",IF(M{row}="Good","Pursue with standard process",'
         f'IF(M{row}="Possible","Offer a lighter or self-serve option",IF(M{row}="Poor","Decline politely or refer","Check red flag before proceeding")))))')
    icp.cell(row=row, column=14).alignment = Alignment(wrap_text=True, vertical="center")
    icp.row_dimensions[row].height = 30
dv_05 = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True,
                       showErrorMessage=True, errorTitle="Rating 0–5", error="Enter a whole number from 0 to 5.")
icp.add_data_validation(dv_05)
dv_05.add(f"D{R_FIRST}:J{R_LAST}")
dv_yn = DataValidation(type="list", formula1="=Lists!$E$2:$E$3", allow_blank=True)
icp.add_data_validation(dv_yn)
dv_yn.add(f"K{R_FIRST}:K{R_LAST}")
dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                         showErrorMessage=True, errorTitle="Date", error="Enter a valid date.")
icp.add_data_validation(dv_date)
dv_date.add(f"C{R_FIRST}:C{R_LAST}")
rng = f"M{R_FIRST}:M{R_LAST}"
icp.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Ideal"'], fill=fill(DARK), font=Font(name=FONT, bold=True, color=WHITE)))
icp.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Good"'], fill=fill("B7D7C2")))
icp.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Possible"'], fill=fill("FBE7B0")))
icp.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(M{R_FIRST}="Poor",M{R_FIRST}="Red flag — review")'], fill=fill("F4D6CC")))
icp.freeze_panes = f"D{R_FIRST}"
page(icp)

# ---------------------------------------------------------------- 5. Interview Log
il = wb.create_sheet("Interview Log", 4)
title(il, "INTERVIEW LOG", "One row per real customer conversation. Use codes, not names or phone numbers. Questions: prompt P-STR-009.", 12)
cols = ["Date", "Customer code", "Persona", "Situation", "Trigger (why now?)", "Pain quotes (exact words)",
        "Current solution", "Budget signal", "Objections / doubts", "Next step", "Pain added to scoring?", "Notes"]
widths(il, {"A": 13, "B": 13, "C": 14, "D": 30, "E": 24, "F": 40, "G": 24, "H": 16, "I": 28, "J": 22, "K": 13, "L": 24})
header(il, 4, cols)
I_FIRST, I_LAST = 5, 104
for i in range(100):
    row = I_FIRST + i
    for col in range(1, 13):
        inp(il.cell(row=row, column=col))
    il.cell(row=row, column=1).number_format = "dd-mmm-yyyy"
    il.row_dimensions[row].height = 30
ex = ["=DATE(2026,9,14)", "Example C-014", "Persona 1", "Couple, both working; 2BHK in Wakad; part-time maid for 2 years",
      "Guests arriving for a family function", "'She does the floor and dishes but the bathroom tiles are never really clean.' · 'I hate checking her work.'",
      "Part-time maid + deep clean before Diwali", "Pays for help already", "'Will it be the same team each time?'",
      "Send Care Plan one-pager", "Yes", "Interested in weekday morning slot"]
for col, v in enumerate(ex, start=1):
    il.cell(row=I_FIRST, column=col, value=v)
il.cell(row=I_FIRST, column=1).number_format = "dd-mmm-yyyy"
dv_p2 = DataValidation(type="list", formula1="=Lists!$A$2:$A$4", allow_blank=True)
il.add_data_validation(dv_p2)
dv_p2.add(f"C{I_FIRST}:C{I_LAST}")
dv_b = DataValidation(type="list", formula1="=Lists!$D$2:$D$6", allow_blank=True)
il.add_data_validation(dv_b)
dv_b.add(f"H{I_FIRST}:H{I_LAST}")
dv_y2 = DataValidation(type="list", formula1="=Lists!$E$2:$E$3", allow_blank=True)
il.add_data_validation(dv_y2)
dv_y2.add(f"K{I_FIRST}:K{I_LAST}")
dv_d2 = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                       showErrorMessage=True, errorTitle="Date", error="Enter a valid date.")
il.add_data_validation(dv_d2)
dv_d2.add(f"A{I_FIRST}:A{I_LAST}")
il["N4"] = "Conversations logged"
il["N4"].font = f(10, True, DARK)
calc(il["N5"], f"=COUNTA(B{I_FIRST}:B{I_LAST})", "0")
il["N7"] = "Pains added to scoring"
il["N7"].font = f(10, True, DARK)
calc(il["N8"], f'=COUNTIF(K{I_FIRST}:K{I_LAST},"Yes")', "0")
il.column_dimensions["M"].width = 3
il.column_dimensions["N"].width = 22
il.freeze_panes = f"C{I_FIRST}"
il.auto_filter.ref = f"A4:L{I_LAST}"
page(il)

# ---------------------------------------------------------------- Lists content
for col, lab in zip("ABCDE", ["Persona", "Source type", "Persona status", "Budget signal", "Yes / No"]):
    c = lists[f"{col}1"]
    c.value = lab
    c.font = f(10, True, WHITE)
    c.fill = fill(SECOND)
    c.border = BORDER
list_values = {
    "A": ["Persona 1", "Persona 2", "Persona 3"],
    "B": ["Interview", "Enquiry / chat", "Email", "Sales call", "Own review", "Competitor review", "Online community", "Other"],
    "C": ["Draft", "Evidence-based", "Validated"],
    "D": ["Pays for help already", "Has budget, not spending", "Price-sensitive", "No budget", "Unknown"],
    "E": ["Yes", "No"],
}
for col, vals in list_values.items():
    for i, v in enumerate(vals):
        lists[f"{col}{i + 2}"] = v
        lists[f"{col}{i + 2}"].font = f()
widths(lists, {"A": 16, "B": 20, "C": 18, "D": 26, "E": 10, "F": 3, "G": 60})
lists["G1"] = "Note"
lists["G1"].font = f(10, True, DARK)
lists["G2"] = ("Persona dropdowns use 'Persona 1/2/3' so they stay stable when you rename personas. "
               "Persona names are set in the Persona Builder. You may add items to columns B and D; "
               "extend the dropdown ranges if you add more rows.")
lists["G2"].alignment = WRAP
lists["G2"].font = f()

wb.move_sheet("Lists", offset=len(wb.sheetnames))
for ws_ in wb.worksheets:
    ws_.sheet_view.showGridLines = False
wb["Start Here"].sheet_properties.tabColor = GOLD
for name in ("Persona Builder", "Pain Point Scoring", "ICP Scorecard", "Interview Log"):
    wb[name].sheet_properties.tabColor = DARK
wb["Lists"].sheet_properties.tabColor = "8A8A8A"
wb.properties.title = "Customer Persona Builder — AI Business Operating System 2026"
wb.properties.creator = "AQVANI.SHOP"
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
